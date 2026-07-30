import base64
import csv
import io
import json
import os
import re
import sys
import time
import unicodedata
from datetime import date, datetime
from pathlib import Path
from typing import Any, Optional
from zipfile import BadZipFile

from openpyxl import load_workbook

CURRENT_DIR = Path(__file__).resolve().parent
if str(CURRENT_DIR) not in sys.path:
    sys.path.insert(0, str(CURRENT_DIR))

try:
    from .matching import SearchRecord, TextCandidateIndex, normalize_text
    from .medical_fees import load_medical_fees_index
    from .medical_fees_excel import load_medical_fees_excel_index
    from .prompts import SYSTEM_PROMPT, USER_PROMPT
    from .schemas import (
        AgentData,
        AgentInput,
        AgentResponse,
        CandidateCode,
        HomologatedItem,
        ProcessingSummary,
        ServiceRow,
    )
except ImportError:
    from matching import SearchRecord, TextCandidateIndex, normalize_text
    from medical_fees import load_medical_fees_index
    from medical_fees_excel import load_medical_fees_excel_index
    from prompts import SYSTEM_PROMPT, USER_PROMPT
    from schemas import (
        AgentData,
        AgentInput,
        AgentResponse,
        CandidateCode,
        HomologatedItem,
        ProcessingSummary,
        ServiceRow,
    )


MODEL_NAME = os.getenv("PROVIDER_SYNC_HOMOLOGADOR_MODEL", "gpt-5.6-luna")
SUPPORTED_FILE_SUFFIXES = {".xlsx", ".xlsm", ".csv", ".tsv", ".txt"}
REQUIRED_COLUMNS = {"codigoServicio", "nombreServicio"}
PDF_FALLBACK_CONFIDENCE = 80

HEADER_ALIASES = {
    "ROW NUMBER": "rowNumber",
    "FILA": "rowNumber",
    "TIPO SERVICIO": "tipoServicio",
    "TIPO DE SERVICIO": "tipoServicio",
    "CODIGO SERVICIO": "codigoServicio",
    "CODIGO ESTUDIO": "codigoServicio",
    "CODIGO PROCEDIMIENTO": "codigoServicio",
    "CODIGO ASIGNADO RAIZA": "codigoServicio",
    "ID PROCEDIMIENTO PROVEEDOR": "codigoServicio",
    "NOMBRE SERVICIO": "nombreServicio",
    "NOMBRE ESTUDIO": "nombreServicio",
    "NOMBRE PROCEDIMIENTO": "nombreServicio",
    "NOMBRE ASIGNADO RAIZA": "nombreServicio",
    "DESCRIPCION": "nombreServicio",
    "MONEDA": "moneda",
    "PRECIO": "precio",
    "PRECIO NEGOCIADO": "precio",
    "TARIFA": "precio",
    "FECHA INICIO": "fechaInicio",
    "FECHA DE INICIO": "fechaInicio",
    "VIGENCIA INICIO": "fechaInicio",
    "OBSERVACION": "observaciones",
    "OBSERVACIONES": "observaciones",
    "TIPO PROVEEDOR": "tipoProveedor",
    "TIPO DE PROVEEDOR": "tipoProveedor",
    "CODIGO PROVEEDOR": "codigoProveedor",
    "ID PROVEEDOR": "codigoProveedor",
    "NOMBRE PROVEEDOR": "nombreProveedor",
    "PROVEEDOR": "nombreProveedor",
}


def build_client() -> Any:
    from openai import OpenAI

    api_key = os.getenv("OPENAI_API_KEY")
    if not api_key:
        raise RuntimeError("No se encontro OPENAI_API_KEY en variables de entorno.")
    return OpenAI(api_key=api_key)


def normalize_header(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(character for character in text if not unicodedata.combining(character))
    text = re.sub(r"[^A-Za-z0-9]+", " ", text).upper()
    return re.sub(r"\s+", " ", text).strip()


def canonical_header(value: Any) -> str:
    original = str(value or "").strip()
    return HEADER_ALIASES.get(normalize_header(value), original)


def normalize_row_keys(row: dict[str, Any], row_number: int) -> dict[str, Any]:
    normalized = {
        canonical_header(key): value
        for key, value in row.items()
        if str(key or "").strip()
    }
    normalized["rowNumber"] = normalized.get("rowNumber") or row_number
    return normalized


def detect_delimiter(text: str, default: str = ",") -> str:
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
        return dialect.delimiter
    except csv.Error:
        first_line = next((line for line in text.splitlines() if line.strip()), "")
        if "\t" in first_line:
            return "\t"
        if ";" in first_line:
            return ";"
        if "|" in first_line:
            return "|"
        return default


def matrix_to_rows(matrix: list[list[Any]]) -> list[dict[str, Any]]:
    non_empty_rows = [
        (index, values)
        for index, values in enumerate(matrix, start=1)
        if any(value not in (None, "") for value in values)
    ]
    if not non_empty_rows:
        return []

    header_position = 0
    best_score = -1
    for position, (_, values) in enumerate(non_empty_rows[:30]):
        canonical = {canonical_header(value) for value in values if value not in (None, "")}
        score = len(canonical & REQUIRED_COLUMNS)
        if score > best_score:
            header_position = position
            best_score = score
        if REQUIRED_COLUMNS.issubset(canonical):
            header_position = position
            break

    _, header_values = non_empty_rows[header_position]
    headers = [canonical_header(value) for value in header_values]
    rows: list[dict[str, Any]] = []
    for sheet_row_number, values in non_empty_rows[header_position + 1:]:
        item = {
            headers[index]: values[index] if index < len(values) else None
            for index in range(len(headers))
            if headers[index]
        }
        if any(value not in (None, "") for value in item.values()):
            rows.append(normalize_row_keys(item, sheet_row_number))
    return rows


def load_rows_from_delimited_text(text: str, delimiter: Optional[str] = None) -> list[dict[str, Any]]:
    selected_delimiter = delimiter or detect_delimiter(text)
    matrix = list(csv.reader(io.StringIO(text), delimiter=selected_delimiter))
    return matrix_to_rows(matrix)


def load_rows_from_excel_bytes(file_bytes: bytes) -> list[dict[str, Any]]:
    workbook = load_workbook(filename=io.BytesIO(file_bytes), read_only=True, data_only=True)
    sheet = workbook.active
    matrix = [list(values) for values in sheet.iter_rows(values_only=True)]
    return matrix_to_rows(matrix)


def load_rows_from_file(file_path: str) -> list[dict[str, Any]]:
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"No existe el archivo: {file_path}")

    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_FILE_SUFFIXES:
        raise ValueError(
            f"Formato no soportado: {suffix}. Usa uno de {sorted(SUPPORTED_FILE_SUFFIXES)}."
        )

    if suffix in {".csv", ".tsv", ".txt"}:
        default_delimiter = "\t" if suffix == ".tsv" else ","
        return load_rows_from_delimited_text(
            path.read_text(encoding="utf-8-sig"),
            delimiter=detect_delimiter(path.read_text(encoding="utf-8-sig"), default_delimiter),
        )

    return load_rows_from_excel_bytes(path.read_bytes())


def decode_base64_file(payload: AgentInput) -> list[dict[str, Any]]:
    if not payload.excelBase64:
        return []

    try:
        file_bytes = base64.b64decode(payload.excelBase64, validate=True)
    except Exception as exc:
        raise ValueError("excelBase64 no contiene un archivo valido en base64.") from exc

    file_name = (payload.fileName or "").lower()
    if file_name.endswith((".csv", ".tsv", ".txt")):
        text = file_bytes.decode("utf-8-sig")
        default_delimiter = "\t" if file_name.endswith(".tsv") else ","
        return load_rows_from_delimited_text(
            text, delimiter=detect_delimiter(text, default_delimiter)
        )

    try:
        return load_rows_from_excel_bytes(file_bytes)
    except BadZipFile as exc:
        raise ValueError(
            "No se pudo leer el archivo base64 como Excel. Envia fileName con la extension correcta."
        ) from exc


def collect_rows(payload: AgentInput) -> list[ServiceRow]:
    raw_rows: list[dict[str, Any]] = []

    if payload.rows:
        raw_rows.extend(payload.rows)
    if payload.filePath:
        raw_rows.extend(load_rows_from_file(payload.filePath))
    if payload.csvText:
        raw_rows.extend(load_rows_from_delimited_text(payload.csvText))
    if payload.excelBase64:
        raw_rows.extend(decode_base64_file(payload))

    if not raw_rows:
        raise ValueError("No se encontraron filas para procesar.")

    validated_rows: list[ServiceRow] = []
    for index, row in enumerate(raw_rows, start=1):
        normalized = normalize_row_keys(row, index)
        validated_rows.append(ServiceRow.model_validate(normalized))
    return validated_rows


def batch_rows(rows: list[ServiceRow], batch_size: int) -> list[list[ServiceRow]]:
    return [rows[index:index + batch_size] for index in range(0, len(rows), batch_size)]


def build_cpt_index(payload: AgentInput) -> TextCandidateIndex:
    records = [
        SearchRecord(
            codigo=item.codigoCpt,
            nombre=item.nombreCpt,
            fuente="catalogo_cpt",
            id_cpt_product=item.idCptProduct,
            tipo_procedimiento=item.tipoProcedimientoNombre,
        )
        for item in payload.catalogoCpt
        if str(item.estado) != "0"
    ]
    return TextCandidateIndex(records)


def merge_candidates(
    *candidate_groups: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    merged: list[dict[str, Any]] = []
    seen: set[tuple[Any, ...]] = set()
    for candidates in candidate_groups:
        for candidate in candidates:
            key = (
                candidate.get("fuente"),
                str(candidate.get("codigo") or ""),
                candidate.get("categoria"),
                candidate.get("paginaMedicalFees"),
                candidate.get("idCptProduct"),
            )
            if key in seen:
                continue
            seen.add(key)
            merged.append(candidate)
    return merged


def enrich_rows(
    rows: list[ServiceRow],
    cpt_index: TextCandidateIndex,
    medical_fees_excel_index: TextCandidateIndex,
    medical_fees_pdf_index: Optional[TextCandidateIndex],
    max_candidates: int,
) -> list[dict[str, Any]]:
    enriched: list[dict[str, Any]] = []
    for row in rows:
        item = row.model_dump(mode="json", exclude_none=True)
        item["candidatosCatalogoCpt"] = cpt_index.search(
            row.codigoServicio, row.nombreServicio, max_candidates
        )
        item["candidatosMedicalFeesExcel"] = medical_fees_excel_index.search(
            row.codigoServicio, row.nombreServicio, max_candidates
        )
        item["candidatosMedicalFeesPdf"] = (
            medical_fees_pdf_index.search(
                row.codigoServicio, row.nombreServicio, max_candidates
            )
            if medical_fees_pdf_index is not None
            else []
        )
        item["busquedaPdfMedicalFees"] = medical_fees_pdf_index is not None
        enriched.append(item)
    return enriched


def build_context(
    payload: AgentInput,
    rows: list[ServiceRow],
    medical_fees_excel_index: TextCandidateIndex,
    medical_fees_pdf_index: Optional[TextCandidateIndex],
    etapa: str,
) -> dict[str, Any]:
    first = rows[0]
    return {
        "archivo": payload.fileName or payload.filePath,
        "totalFilas": len(rows),
        "tipoProveedor": first.tipoProveedor,
        "codigoProveedor": first.codigoProveedor,
        "nombreProveedor": first.nombreProveedor,
        "metadata": payload.metadata or {},
        "catalogoCptRecibido": len(payload.catalogoCpt),
        "codigosMedicalFeesExcelIndexados": len(
            medical_fees_excel_index.records
        ),
        "codigosMedicalFeesPdfIndexados": (
            len(medical_fees_pdf_index.records)
            if medical_fees_pdf_index is not None
            else 0
        ),
        "etapaHomologacion": etapa,
        "umbralFallbackPdf": PDF_FALLBACK_CONFIDENCE,
        "catalogoTiposServicio": [
            item.model_dump(mode="json", exclude_none=True)
            for item in payload.catalogoTiposServicio
        ],
        "prioridadFuentes": [
            "catalogo_cpt",
            "medical_fees_excel",
            "medical_fees_pdf_solo_fallback",
            "sin_coincidencia",
        ],
    }


def extract_json_from_text(output_text: str) -> dict[str, Any]:
    cleaned = output_text.strip()
    if cleaned.startswith("```"):
        lines = cleaned.splitlines()
        if lines and lines[0].startswith("```"):
            lines = lines[1:]
        if lines and lines[-1].strip() == "```":
            lines = lines[:-1]
        cleaned = "\n".join(lines).strip()

    start = cleaned.find("{")
    end = cleaned.rfind("}")
    if start == -1 or end == -1 or end < start:
        raise ValueError("El modelo no devolvio un objeto JSON valido.")
    return json.loads(cleaned[start:end + 1])


def run_agent_for_batch(
    batch: list[ServiceRow],
    context: dict[str, Any],
    cpt_index: TextCandidateIndex,
    medical_fees_excel_index: TextCandidateIndex,
    medical_fees_pdf_index: Optional[TextCandidateIndex],
    max_candidates: int,
    client: Any,
    enable_web_search: bool,
) -> tuple[list[HomologatedItem], Optional[int], dict[int, list[dict[str, Any]]]]:
    enriched_rows = enrich_rows(
        batch,
        cpt_index,
        medical_fees_excel_index,
        medical_fees_pdf_index,
        max_candidates,
    )
    candidate_map = {
        int(item["rowNumber"]): merge_candidates(
            item["candidatosCatalogoCpt"],
            item["candidatosMedicalFeesExcel"],
            item["candidatosMedicalFeesPdf"],
        )
        for item in enriched_rows
    }
    rendered_prompt = USER_PROMPT.format(
        context_json=json.dumps(context, ensure_ascii=False, indent=2),
        rows_json=json.dumps(enriched_rows, ensure_ascii=False, indent=2),
    )

    request_args: dict[str, Any] = {
        "model": MODEL_NAME,
        "reasoning": {"effort": "medium"},
        "input": [
            {"role": "system", "content": SYSTEM_PROMPT},
            {"role": "user", "content": rendered_prompt},
        ],
    }
    if enable_web_search:
        request_args["tools"] = [{"type": "web_search_preview"}]

    response = client.responses.create(**request_args)
    parsed = extract_json_from_text((response.output_text or "").strip())
    items = []
    for raw_item in parsed.get("items", []):
        if not isinstance(raw_item, dict):
            continue
        normalized_item = dict(raw_item)
        safe_translations = []
        for translation in normalized_item.get("traduccionesMedicalFees", []):
            if not isinstance(translation, dict):
                continue
            code = str(translation.get("codigo") or "").strip()
            spanish_name = str(
                translation.get("nombreEspanol") or ""
            ).strip()
            if code and spanish_name:
                safe_translations.append(
                    {
                        "codigo": code,
                        "nombreEspanol": spanish_name,
                    }
                )
        normalized_item["traduccionesMedicalFees"] = safe_translations
        # Los candidatos verificados se agregan localmente; nunca se confia en los
        # candidatos que el modelo pueda reformatear o inventar.
        normalized_item["codigosCandidatos"] = []
        items.append(HomologatedItem.model_validate(normalized_item))
    total_tokens = getattr(getattr(response, "usage", None), "total_tokens", None)
    return items, total_tokens, candidate_map


def confidence_label(percentage: int) -> str:
    if percentage >= 80:
        return "alta"
    if percentage >= 65:
        return "media"
    return "baja"


def to_json_date(value: Any) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)


def candidate_models(
    candidates: list[dict[str, Any]],
    translations: Optional[dict[str, str]] = None,
) -> list[CandidateCode]:
    translated_names = translations or {}
    return [
        CandidateCode(
            fuente=candidate["fuente"],
            idCptProduct=candidate.get("idCptProduct"),
            codigo=str(candidate["codigo"]),
            nombre=(
                translated_names.get(str(candidate["codigo"]), candidate["nombre"])
                if candidate["fuente"] == "medical_fees"
                else candidate["nombre"]
            ),
            nombreOriginal=(
                candidate["nombre"]
                if candidate["fuente"] == "medical_fees"
                else None
            ),
            tipoProcedimiento=candidate.get("tipoProcedimiento"),
            paginaMedicalFees=candidate.get("paginaMedicalFees"),
            categoria=candidate.get("categoria"),
            scorePreliminar=candidate.get("scorePreliminar"),
            motivo=(
                f"Coincidencia preliminar de texto y codigo: "
                f"{candidate.get('scorePreliminar', 0)}%."
            ),
        )
        for candidate in candidates
    ]


def find_verified_candidate(
    item: HomologatedItem,
    candidates: list[dict[str, Any]],
) -> Optional[dict[str, Any]]:
    matches = [
        candidate
        for candidate in candidates
        if candidate["fuente"] == item.fuenteAsociacion
        and str(candidate["codigo"]) == str(item.codigoCpt or "")
    ]
    if not matches:
        return None
    if item.categoria:
        return next(
            (
                candidate
                for candidate in matches
                if candidate.get("categoria") == item.categoria
            ),
            None,
        )
    if item.paginaMedicalFees is not None:
        return next(
            (
                candidate
                for candidate in matches
                if candidate.get("paginaMedicalFees")
                == item.paginaMedicalFees
            ),
            None,
        )
    return matches[0]


def normalize_agent_item(
    item: HomologatedItem,
    row: ServiceRow,
    candidates: list[dict[str, Any]],
    payload: AgentInput,
) -> HomologatedItem:
    selected = find_verified_candidate(item, candidates)

    if item.fuenteAsociacion != "sin_coincidencia" and selected is None:
        return fallback_item(
            row,
            candidates,
            "El modelo propuso un codigo que no estaba en los candidatos verificados.",
        )

    service_type_id = item.tipoServicioId
    service_type_name = item.tipoServicioNombre
    valid_service_type = None
    for service_type in payload.catalogoTiposServicio:
        same_id = (
            service_type_id is not None
            and str(service_type.idTipoServicio) == str(service_type_id)
        )
        same_name = (
            service_type_name
            and normalize_text(service_type.nombreTipoServicio)
            == normalize_text(service_type_name)
        )
        if same_id or same_name:
            valid_service_type = service_type
            break

    medical_fees_translations = {
        str(translation.codigo): translation.nombreEspanol
        for translation in item.traduccionesMedicalFees
    }
    if (
        item.fuenteAsociacion == "medical_fees"
        and item.codigoCpt
        and item.nombreCpt
    ):
        medical_fees_translations.setdefault(str(item.codigoCpt), item.nombreCpt)

    updates: dict[str, Any] = {
        "codigoServicio": row.codigoServicio,
        "nombreServicioOriginal": row.nombreServicio,
        "precio": row.precio,
        "fechaInicio": to_json_date(row.fechaInicio),
        "observacionesArchivo": row.observaciones,
        "confianza": confidence_label(item.confianzaPorcentaje),
        "requiereRevisionHumana": item.confianzaPorcentaje < 80,
        "codigosCandidatos": candidate_models(
            candidates,
            medical_fees_translations,
        ),
    }

    if valid_service_type:
        updates["tipoServicioId"] = valid_service_type.idTipoServicio
        updates["tipoServicioNombre"] = valid_service_type.nombreTipoServicio
    else:
        updates["tipoServicioId"] = None
        updates["tipoServicioNombre"] = None
        updates["requiereRevisionHumana"] = True

    if selected:
        translated_cpt_name = item.nombreCpt
        updates["codigoCpt"] = selected["codigo"]
        updates["nombreCptOriginal"] = selected["nombre"]
        updates["idCptProduct"] = selected.get("idCptProduct")
        updates["paginaMedicalFees"] = selected.get("paginaMedicalFees")
        updates["categoria"] = selected.get("categoria")
        updates["codigoMedicalFees"] = selected["codigo"]

        if selected["fuente"] == "catalogo_cpt":
            updates["nombreCpt"] = selected["nombre"]
            updates["estadoAsociacion"] = "asociado"
            updates["fuenteAsociacion"] = "catalogo_cpt"
        else:
            updates["nombreCpt"] = (
                selected["nombre"]
                if selected.get("categoria")
                else translated_cpt_name or selected["nombre"]
            )
            updates["estadoAsociacion"] = "cpt_nuevo_sugerido"
            updates["fuenteAsociacion"] = "medical_fees"
    else:
        updates.update(
            {
                "estadoAsociacion": "sin_asociacion",
                "fuenteAsociacion": "sin_coincidencia",
                "idCptProduct": None,
                "codigoCpt": None,
                "nombreCpt": None,
                "nombreCptOriginal": None,
                "codigoMedicalFees": None,
                "paginaMedicalFees": None,
                "categoria": None,
                "requiereRevisionHumana": True,
            }
        )

    if item.confianzaPorcentaje < 80:
        updates["requiereRevisionHumana"] = True
    return item.model_copy(update=updates)


def fallback_item(
    row: ServiceRow,
    candidates: list[dict[str, Any]],
    observation: str,
) -> HomologatedItem:
    return HomologatedItem(
        rowNumber=row.rowNumber or 0,
        codigoServicio=row.codigoServicio,
        nombreServicioOriginal=row.nombreServicio,
        precio=row.precio,
        fechaInicio=to_json_date(row.fechaInicio),
        observacionesArchivo=row.observaciones,
        estadoAsociacion="sin_asociacion",
        fuenteAsociacion="sin_coincidencia",
        confianzaPorcentaje=0,
        confianza="baja",
        requiereRevisionHumana=True,
        motivo=observation,
        codigosCandidatos=candidate_models(candidates),
        fuentesWeb=[],
        observaciones=observation,
    )


def normalize_stage_results(
    rows: list[ServiceRow],
    raw_items: list[HomologatedItem],
    candidates_by_row: dict[int, list[dict[str, Any]]],
    payload: AgentInput,
) -> dict[int, HomologatedItem]:
    raw_by_row = {item.rowNumber: item for item in raw_items}
    normalized: dict[int, HomologatedItem] = {}
    for index, row in enumerate(rows, start=1):
        row_number = row.rowNumber or index
        candidates = candidates_by_row.get(row_number, [])
        if row_number in raw_by_row:
            normalized[row_number] = normalize_agent_item(
                raw_by_row[row_number],
                row,
                candidates,
                payload,
            )
        else:
            normalized[row_number] = fallback_item(
                row,
                candidates,
                "El modelo no devolvio homologacion para esta fila.",
            )
    return normalized


def requires_pdf_fallback(item: HomologatedItem) -> bool:
    return (
        item.estadoAsociacion == "sin_asociacion"
        or item.confianzaPorcentaje < PDF_FALLBACK_CONFIDENCE
    )


def result_priority(item: HomologatedItem) -> tuple[int, int, int]:
    association_rank = 0 if item.estadoAsociacion == "sin_asociacion" else 1
    if item.fuenteAsociacion == "catalogo_cpt":
        source_rank = 3
    elif item.fuenteAsociacion == "medical_fees" and item.categoria:
        source_rank = 2
    elif item.fuenteAsociacion == "medical_fees":
        source_rank = 1
    else:
        source_rank = 0
    return association_rank, item.confianzaPorcentaje, source_rank


def choose_fallback_result(
    initial: HomologatedItem,
    retried: HomologatedItem,
) -> HomologatedItem:
    if result_priority(retried) > result_priority(initial):
        return retried
    return initial.model_copy(
        update={"codigosCandidatos": retried.codigosCandidatos}
    )


def process_request_payload(
    raw_payload: dict[str, Any],
    client: Optional[Any] = None,
    medical_fees_index: Optional[TextCandidateIndex] = None,
    medical_fees_excel_index: Optional[TextCandidateIndex] = None,
) -> AgentResponse:
    started = time.monotonic()
    payload = AgentInput.model_validate(raw_payload)
    rows = collect_rows(payload)
    cpt_index = build_cpt_index(payload)
    excel_index = (
        medical_fees_excel_index
        or load_medical_fees_excel_index()
    )

    if client is None:
        client = build_client()

    initial_context = build_context(
        payload,
        rows,
        excel_index,
        None,
        "excel_primero",
    )
    initial_raw_items: list[HomologatedItem] = []
    initial_candidates: dict[int, list[dict[str, Any]]] = {}
    total_tokens = 0

    for batch in batch_rows(rows, payload.batchSize):
        items, batch_tokens, candidate_map = run_agent_for_batch(
            batch=batch,
            context=initial_context,
            cpt_index=cpt_index,
            medical_fees_excel_index=excel_index,
            medical_fees_pdf_index=None,
            max_candidates=payload.maxCandidates,
            client=client,
            enable_web_search=payload.enableWebSearch,
        )
        initial_raw_items.extend(items)
        initial_candidates.update(candidate_map)
        if batch_tokens:
            total_tokens += batch_tokens

    final_by_row = normalize_stage_results(
        rows,
        initial_raw_items,
        initial_candidates,
        payload,
    )
    fallback_rows = [
        row
        for index, row in enumerate(rows, start=1)
        if requires_pdf_fallback(final_by_row[row.rowNumber or index])
    ]

    if fallback_rows:
        pdf_index = medical_fees_index or load_medical_fees_index()
        fallback_context = build_context(
            payload,
            fallback_rows,
            excel_index,
            pdf_index,
            "fallback_pdf",
        )
        fallback_raw_items: list[HomologatedItem] = []
        fallback_candidates: dict[int, list[dict[str, Any]]] = {}
        for batch in batch_rows(fallback_rows, payload.batchSize):
            items, batch_tokens, candidate_map = run_agent_for_batch(
                batch=batch,
                context=fallback_context,
                cpt_index=cpt_index,
                medical_fees_excel_index=excel_index,
                medical_fees_pdf_index=pdf_index,
                max_candidates=payload.maxCandidates,
                client=client,
                enable_web_search=payload.enableWebSearch,
            )
            fallback_raw_items.extend(items)
            fallback_candidates.update(candidate_map)
            if batch_tokens:
                total_tokens += batch_tokens

        retried_by_row = normalize_stage_results(
            fallback_rows,
            fallback_raw_items,
            fallback_candidates,
            payload,
        )
        for row_number, retried in retried_by_row.items():
            final_by_row[row_number] = choose_fallback_result(
                final_by_row[row_number],
                retried,
            )

    ordered_items = [
        final_by_row[row.rowNumber or index]
        for index, row in enumerate(rows, start=1)
    ]

    elapsed = f"{(time.monotonic() - started):.6f}"
    summary = ProcessingSummary(
        totalItems=len(ordered_items),
        itemsAsociadosCatalogo=sum(
            1 for item in ordered_items if item.estadoAsociacion == "asociado"
        ),
        itemsCptNuevoSugerido=sum(
            1 for item in ordered_items if item.estadoAsociacion == "cpt_nuevo_sugerido"
        ),
        itemsSinAsociacion=sum(
            1 for item in ordered_items if item.estadoAsociacion == "sin_asociacion"
        ),
        itemsParaRevision=sum(
            1 for item in ordered_items if item.confianzaPorcentaje < 80
        ),
    )

    first_row = rows[0]
    data = AgentData(
        proveedor=first_row.nombreProveedor,
        codigoProveedor=(
            str(first_row.codigoProveedor)
            if first_row.codigoProveedor is not None
            else None
        ),
        tipoProveedor=(
            str(first_row.tipoProveedor)
            if first_row.tipoProveedor is not None
            else None
        ),
        archivoProcesado=payload.fileName or payload.filePath,
        totalFilasEntrada=len(rows),
        items=ordered_items,
        resumen=summary,
    )
    return AgentResponse(
        modelo=MODEL_NAME,
        tokensUsados=total_tokens or None,
        tiempoRespuestaSegundos=elapsed,
        data=data,
    )


def main() -> None:
    if len(sys.argv) != 2:
        raise SystemExit(
            'Uso: python3 pipeline.py \'{"filePath":"/ruta/catalogo.xlsx",'
            '"catalogoCpt":[],"catalogoTiposServicio":[]}\''
        )

    payload = json.loads(sys.argv[1])
    result = process_request_payload(payload)
    print(
        json.dumps(
            result.model_dump(mode="json", exclude_none=True),
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
