import os
import re
import unicodedata
from functools import lru_cache
from pathlib import Path
from typing import Any, Optional

from openpyxl import load_workbook

try:
    from .matching import SearchRecord, TextCandidateIndex
except ImportError:
    from matching import SearchRecord, TextCandidateIndex


CURRENT_DIR = Path(__file__).resolve().parent
DEFAULT_EXCEL_PATHS = (
    CURRENT_DIR / "CODIGO_MEDICAL_FEES.xlsx",
    CURRENT_DIR / "CODIGO MEDICAL FEES.xlsx",
)
CODE_HEADERS = {
    "CODIGO",
    "CODIGO NUEVO 2022",
    "NUEVA CODIFICACION",
    "NUEVA CODIFICACION MEDICAL FEES",
    "NUEVO CODIGO 2022",
    "NUEVO CODIGO 2022 MEDICAL FEES",
}
DESCRIPTION_HEADERS = {
    "DESCRIPCION",
    "NOMBRE",
    "PROCEDIMIENTO",
}
VALID_CODE = re.compile(r"^\d{5,12}(?:\.\d+)?$")


def normalize_header(value: Any) -> str:
    normalized = unicodedata.normalize("NFKD", str(value or ""))
    normalized = "".join(
        character
        for character in normalized
        if not unicodedata.combining(character)
    )
    normalized = re.sub(r"[^A-Za-z0-9]+", " ", normalized).upper()
    return re.sub(r"\s+", " ", normalized).strip()


def clean_code(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    text = str(value).strip()
    if re.fullmatch(r"\d+\.0", text):
        return text[:-2]
    return text


def find_header_columns(rows: list[tuple[Any, ...]]) -> tuple[int, int, int]:
    for row_index, values in enumerate(rows):
        code_column = None
        description_column = None
        for column_index, value in enumerate(values):
            normalized = normalize_header(value)
            if normalized in CODE_HEADERS:
                code_column = column_index
            if normalized in DESCRIPTION_HEADERS:
                description_column = column_index
        if code_column is not None and description_column is not None:
            return row_index, code_column, description_column
    raise ValueError("No se encontraron columnas de codigo y descripcion.")


def resolve_excel_path() -> Path:
    configured_path = os.getenv("MEDICAL_FEES_EXCEL_PATH")
    if configured_path:
        return Path(configured_path)
    for candidate in DEFAULT_EXCEL_PATHS:
        if candidate.exists():
            return candidate
    return DEFAULT_EXCEL_PATHS[0]


@lru_cache(maxsize=2)
def load_medical_fees_excel_records(
    excel_path_value: Optional[str] = None,
) -> tuple[SearchRecord, ...]:
    excel_path = Path(excel_path_value) if excel_path_value else resolve_excel_path()
    if not excel_path.exists():
        raise FileNotFoundError(
            f"No se encontro el Excel de codigos Medical Fees en {excel_path}."
        )

    records: list[SearchRecord] = []
    seen: set[tuple[str, str, str]] = set()
    workbook = load_workbook(
        filename=excel_path,
        read_only=True,
        data_only=True,
    )
    try:
        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue
            try:
                header_row, code_column, description_column = find_header_columns(
                    rows[:30]
                )
            except ValueError:
                continue

            categoria = sheet.title.strip()
            for values in rows[header_row + 1:]:
                code = clean_code(
                    values[code_column] if code_column < len(values) else None
                )
                description = str(
                    values[description_column]
                    if description_column < len(values)
                    and values[description_column] is not None
                    else ""
                ).strip()
                if not VALID_CODE.fullmatch(code) or not description:
                    continue

                record_key = (categoria, code, description)
                if record_key in seen:
                    continue
                seen.add(record_key)
                records.append(
                    SearchRecord(
                        codigo=code,
                        nombre=description,
                        fuente="medical_fees",
                        categoria=categoria,
                    )
                )
    finally:
        workbook.close()

    if not records:
        raise RuntimeError(
            "No se pudieron indexar codigos desde el Excel Medical Fees."
        )
    return tuple(records)


@lru_cache(maxsize=2)
def load_medical_fees_excel_index(
    excel_path_value: Optional[str] = None,
) -> TextCandidateIndex:
    return TextCandidateIndex(load_medical_fees_excel_records(excel_path_value))
